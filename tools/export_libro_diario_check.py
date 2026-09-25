#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Inspecciona un .xlsx generado por la app y reporta que quedó bien (F75).

Uso:  python tools/export_libro_diario_check.py "<ruta del reporte.xlsx>"

Lo usa la verificación EN VIVO (`tools/verify_export_excel.mjs`) después de apretar «Exportar Excel»
en la app: acá se abre el archivo REAL con datos reales y se comprueba la estructura y los números.
"""
import sys
from openpyxl import load_workbook

if len(sys.argv) != 2:
    print("Uso: export_libro_diario_check.py <reporte.xlsx>", file=sys.stderr)
    sys.exit(1)

wb = load_workbook(sys.argv[1])
print("HOJAS: " + " | ".join(wb.sheetnames))

ESPERADAS = ["Resumen", "Por método de pago", "Pago Móvil", "Ventas", "Servicios", "IVA del período"]
faltan = [h for h in ESPERADAS if h not in wb.sheetnames]
print("FALTAN: " + (", ".join(faltan) if faltan else "(ninguna)"))

# Por método de pago: cuántas filas y si están ordenadas.
ws = wb["Por método de pago"]
metodos = [ws.cell(row=r, column=1).value for r in range(5, ws.max_row + 1)
           if ws.cell(row=r, column=1).value and ws.cell(row=r, column=2).value]
print(f"POR METODO: {len(metodos)} cobros · ordenados={metodos == sorted(metodos, key=lambda m: str(m).lower())}")
print("METODOS: " + " | ".join(sorted({str(m) for m in metodos})))

# Pago Móvil: montos y referencias.
ws = wb["Pago Móvil"]
pm = [(ws.cell(row=r, column=7).value, ws.cell(row=r, column=8).value) for r in range(5, ws.max_row + 1)
      if ws.cell(row=r, column=1).value and ws.cell(row=r, column=8).value is not None]
con_ref = [x for x in pm if x[0] and "sin referencia" not in str(x[0])]
print(f"PAGO MOVIL: {len(pm)} pagos · {len(con_ref)} con referencia")
for ref, monto in pm[:6]:
    print(f"   {ref} -> {monto}")

# IVA del período.
ws = wb["IVA del período"]
print("IVA:")
for r in range(5, ws.max_row + 1):
    tipo, alic = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value
    if tipo in ("Ventas", "Servicios", "TOTALES") and alic is not None or (tipo == "TOTALES"):
        print(f"   {tipo} {alic} · ops={ws.cell(row=r, column=3).value} · base={ws.cell(row=r, column=4).value}"
              f" · IVA={ws.cell(row=r, column=5).value} · total={ws.cell(row=r, column=6).value}")

# Resumen: los totales del período.
ws = wb["Resumen"]
print("RESUMEN:")
for r in range(1, min(ws.max_row, 40) + 1):
    label = str(ws.cell(row=r, column=1).value or "")
    if label.startswith(("Ventas cobradas", "IVA del período", "IVA de ", "Gastos", "Base imponible")):
        print(f"   {label}: {ws.cell(row=r, column=2).value}")
print("OK")
